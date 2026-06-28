from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.constants import FileAssetKind, ImportBatchStatus, PriceDocumentStatus, ProcessingStatus
from app.db.base import Base
from app.db.models import FileAsset, ImportBatch, PriceDocument, PriceItemVersion, ProcessingEvent, Service
from app.db.session import get_db
from app.main import create_app
from app.workers.pipeline import WorkerPipelineService


@pytest.fixture()
def db_session() -> Iterator[Session]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    with session_factory() as session:
        yield session
    Base.metadata.drop_all(engine)


def test_worker_pipeline_processes_batch_and_keeps_failures_isolated(
    tmp_path: Path,
    db_session: Session,
) -> None:
    batch, good_document, bad_document = create_batch_with_documents(tmp_path, db_session)
    service = WorkerPipelineService(db_session)

    def process_now(price_document_id: str) -> None:
        service.process_document(UUID(price_document_id))

    result = service.process_batch(batch.id, enqueue_document=process_now)

    db_session.refresh(good_document)
    db_session.refresh(bad_document)
    db_session.refresh(batch)
    events = db_session.scalars(select(ProcessingEvent).order_by(ProcessingEvent.created_at)).all()

    assert result["documents"] == 2
    assert good_document.status == PriceDocumentStatus.PARSED.value
    assert good_document.progress_percent == 100
    assert good_document.parsed_summary["parser_format"] == "xlsx"
    assert bad_document.status == PriceDocumentStatus.FAILED.value
    assert bad_document.last_error
    assert batch.processed_files == 1
    assert batch.failed_files == 1
    assert batch.status == ImportBatchStatus.FAILED.value
    assert {event.event_type for event in events} >= {
        "batch_processing_started",
        "document_processing_completed",
        "document_processing_failed",
        "batch_documents_enqueued",
    }
    assert any(event.status == ProcessingStatus.FAILED.value for event in events)


def test_worker_pipeline_records_price_items_for_auto_matched_rows(
    tmp_path: Path,
    db_session: Session,
) -> None:
    batch, document, _ = create_batch_with_documents(tmp_path, db_session)
    db_session.add(
        Service(
            import_batch="test",
            source_type="fixture",
            source_hash="blood-test",
            code="A-1",
            name_ru="Blood test",
            normalized_name="blood test",
            warnings=[],
            raw_data={},
        )
    )
    db_session.commit()

    outcome = WorkerPipelineService(db_session).process_document(document.id)

    db_session.refresh(document)
    price_item = db_session.scalars(select(PriceItemVersion)).one()
    assert outcome.status == PriceDocumentStatus.PARSED.value
    assert document.parsed_summary["normalized_rows"] == 1
    assert document.parsed_summary["recorded_price_items"] == 1
    assert document.parsed_summary["auto_matched"] == 1
    assert price_item.service_id is not None
    assert price_item.amount == 1000
    assert batch.id == document.import_batch_id


def test_reprocess_service_resets_and_processes_single_document(
    tmp_path: Path,
    db_session: Session,
) -> None:
    _, document, _ = create_batch_with_documents(tmp_path, db_session)
    service = WorkerPipelineService(db_session)
    service.process_document(document.id)
    db_session.refresh(document)
    assert document.status == PriceDocumentStatus.PARSED.value

    service.reset_for_reprocess(document.id)
    db_session.refresh(document)
    assert document.status == PriceDocumentStatus.PENDING.value
    assert document.progress_percent == 0

    outcome = service.process_document(document.id, force=True)
    db_session.refresh(document)

    assert outcome.status == PriceDocumentStatus.PARSED.value
    assert document.processing_attempts == 2
    assert db_session.scalar(
        select(ProcessingEvent).where(ProcessingEvent.event_type == "document_reprocess_requested")
    )


def test_reprocess_endpoint_enqueues_single_document(
    tmp_path: Path,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, document, _ = create_batch_with_documents(tmp_path, db_session)
    app = create_app()

    def override_get_db() -> Iterator[Session]:
        yield db_session

    class FakeAsyncResult:
        id = "task-123"

    def fake_delay(price_document_id: str, force: bool = False) -> FakeAsyncResult:
        assert price_document_id == str(document.id)
        assert force is True
        return FakeAsyncResult()

    monkeypatch.setattr("app.api.routes.admin.process_document_task.delay", fake_delay)
    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    token = client.post("/admin/login", json={"username": "admin", "password": "admin"}).json()["access_token"]
    client.headers.update({"Authorization": f"Bearer {token}"})

    response = client.post(f"/admin/import/documents/{document.id}/reprocess")

    assert response.status_code == 200
    assert response.json()["task_id"] == "task-123"
    db_session.refresh(document)
    assert document.status == PriceDocumentStatus.PENDING.value
    assert db_session.scalar(
        select(ProcessingEvent).where(ProcessingEvent.event_type == "document_reprocess_requested")
    )
    app.dependency_overrides.clear()


def create_batch_with_documents(tmp_path: Path, db: Session) -> tuple[ImportBatch, PriceDocument, PriceDocument]:
    storage_root = tmp_path / "storage"
    storage_root.mkdir()
    good_path = storage_root / "prices.xlsx"
    bad_path = storage_root / "notes.txt"
    create_workbook(good_path)
    bad_path.write_text("plain text is unsupported", encoding="utf-8")

    batch = ImportBatch(
        source_type="zip",
        status=ImportBatchStatus.PENDING.value,
        original_filename="archive.zip",
        source_path=str(storage_root / "archive.zip"),
        sha256="batch",
        total_files=2,
        warnings=[],
    )
    db.add(batch)
    db.flush()

    good_asset = FileAsset(
        import_batch_id=batch.id,
        asset_kind=FileAssetKind.ARCHIVE_MEMBER.value,
        original_filename="prices.xlsx",
        stored_path=str(good_path),
        sha256="good",
        size_bytes=good_path.stat().st_size,
        extension=".xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        warnings=[],
    )
    bad_asset = FileAsset(
        import_batch_id=batch.id,
        asset_kind=FileAssetKind.ARCHIVE_MEMBER.value,
        original_filename="notes.txt",
        stored_path=str(bad_path),
        sha256="bad",
        size_bytes=bad_path.stat().st_size,
        extension=".txt",
        mime_type="text/plain",
        warnings=[],
    )
    db.add_all([good_asset, bad_asset])
    db.flush()

    good_document = PriceDocument(
        import_batch_id=batch.id,
        file_asset_id=good_asset.id,
        status=PriceDocumentStatus.PENDING.value,
        detected_type="xlsx",
        warnings=[],
    )
    bad_document = PriceDocument(
        import_batch_id=batch.id,
        file_asset_id=bad_asset.id,
        status=PriceDocumentStatus.PENDING.value,
        detected_type=None,
        warnings=[],
    )
    db.add_all([good_document, bad_document])
    db.commit()
    return batch, good_document, bad_document


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["Code", "Service", "Price"])
    sheet.append(["A-1", "Blood test", 1000])
    workbook.save(path)
